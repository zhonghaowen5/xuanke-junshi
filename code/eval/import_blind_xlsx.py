"""把队友收集的盲评 Excel 导入成标准 blind.csv，并生成统计报告。

用法（在项目 code 目录下执行）：
    python eval\\import_blind_xlsx.py --xlsx "C:\\path\\to\\DeepSeek问题_三分表(1).xlsx"

输出：
    eval/blind.csv            140 条标准盲评数据（含 source_ai 列）
    eval/盲评统计结果.md       三维度均分、分布、来源猜测、回填用表格

说明：
- 兼容 3 个 sheet（DeepSeek / 千问 / 元宝），sheet 名写入 source_ai 列。
- 分数列可能是 int / float / str，统一转 float；空值跳过。
"""
import argparse
import csv
import sys
from collections import Counter
from pathlib import Path

FIELDS = [
    "id", "source_ai", "student", "major", "question",
    "answer", "cites", "rejected", "acc", "use", "trust",
    "comment", "guess",
]


def to_num(v):
    if v is None:
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def load_rows(xlsx: Path):
    """读取所有 sheet，返回 dict 列表（带 source_ai）。"""
    try:
        import openpyxl
    except ImportError:
        print("需要 openpyxl：pip install openpyxl", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    rows = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        it = ws.iter_rows(values_only=True)
        try:
            hdr = [str(c).strip() if c is not None else "" for c in next(it)]
        except StopIteration:
            continue
        for r in it:
            if r is None or r[0] is None:
                continue
            d = dict(zip(hdr, r))
            d["_sheet"] = sn
            rows.append(d)
    return rows


def normalize(rows):
    """映射成标准字段名。"""
    out = []
    for i, d in enumerate(rows, 1):
        rec = {
            "id": d.get("id") or i,
            "source_ai": d.get("_sheet") or d.get("来源AI") or "",
            "student": d.get("student") or d.get("同学") or "",
            "major": d.get("major") or d.get("专业") or "",
            "question": d.get("question") or d.get("问题") or "",
            "answer": d.get("answer") or d.get("回答") or "",
            "cites": d.get("cites") or d.get("出处") or "",
            "rejected": d.get("rejected") or d.get("是否拒答") or "",
            "acc": d.get("acc") or d.get("准确性") or "",
            "use": d.get("use") or d.get("有用性") or "",
            "trust": d.get("trust") or d.get("可信度") or "",
            "comment": d.get("comment") or d.get("评语") or "",
            "guess": d.get("guess") or d.get("来源猜测") or "",
        }
        out.append(rec)
    return out


def stat(rows):
    """计算统计指标，返回 dict。"""
    n = len(rows)
    dims = {"acc": "准确性", "use": "有用性", "trust": "可信度"}
    res = {"n": n}

    allv = []
    for key, label in dims.items():
        vals = [to_num(r[key]) for r in rows]
        vals = [v for v in vals if v is not None]
        if not vals:
            continue
        mean = sum(vals) / len(vals)
        full = 100 * sum(1 for v in vals if v >= 5) / len(vals)
        ge4 = 100 * sum(1 for v in vals if v >= 4) / len(vals)
        res[key] = {
            "label": label, "mean": mean, "n": len(vals),
            "full_pct": full, "ge4_pct": ge4,
            "dist": Counter(int(v) for v in vals),
        }
        allv.extend(vals)

    res["overall"] = sum(allv) / len(allv) if allv else 0
    res["overall_n"] = len(allv)

    # 拒答
    rejected = [r for r in rows if str(r["rejected"]).strip() == "是"]
    res["rejected_n"] = len(rejected)
    res["rejected_pct"] = 100 * len(rejected) / n if n else 0

    # 来源猜测
    res["guess"] = Counter(str(r["guess"]).strip() for r in rows if str(r["guess"]).strip())

    # 覆盖
    res["majors"] = len(set(r["major"] for r in rows if r["major"]))
    res["students"] = len(set(r["student"] for r in rows if r["student"]))

    # 按来源分组
    by_src = {}
    for r in rows:
        s = r["source_ai"]
        by_src.setdefault(s, []).append(r)
    res["by_source"] = {}
    for s, rs in by_src.items():
        vs = []
        for r in rs:
            for k in dims:
                v = to_num(r[k])
                if v is not None:
                    vs.append(v)
        res["by_source"][s] = {
            "n": len(rs),
            "mean": sum(vs) / len(vs) if vs else 0,
            "students": len(set(x["student"] for x in rs)),
        }
    return res


def fmt(x, nd=2):
    return f"{x:.{nd}f}"


def build_md(rows, res):
    """生成回填用的 Markdown 统计报告。"""
    L = []
    L.append("# 实验三 · 同学盲评统计结果\n")
    L.append(f"- 有效问答：**{res['n']} 条**")
    L.append(f"- 覆盖专业：**{res['majors']} 个**")
    L.append(f"- 参与同学编号：**{res['students']} 个**\n")

    L.append("## 一、三维度评分\n")
    L.append("| 维度 | 样本数 | 均分（满分 5） | 满分 5 占比 | ≥4 分占比 |")
    L.append("|---|---|---|---|---|")
    for k in ["acc", "use", "trust"]:
        if k not in res:
            continue
        d = res[k]
        L.append(f"| {d['label']} | {d['n']} | **{fmt(d['mean'])}** | "
                 f"{fmt(d['full_pct'],1)}% | {fmt(d['ge4_pct'],1)}% |")
    L.append(f"\n**综合均分：{fmt(res['overall'])} / 5**（共 {res['overall_n']} 个评分点）\n")

    L.append("## 二、评分分布\n")
    L.append("| 维度 | 5分 | 4分 | 3分 | ≤2分 |")
    L.append("|---|---|---|---|---|")
    for k in ["acc", "use", "trust"]:
        if k not in res:
            continue
        d = res[k]
        dist = d["dist"]
        low = sum(v for kk, v in dist.items() if kk <= 2)
        L.append(f"| {d['label']} | {dist.get(5,0)} | {dist.get(4,0)} | "
                 f"{dist.get(3,0)} | {low} |")

    L.append("\n## 三、拒答与来源猜测\n")
    L.append(f"- 拒答（系统主动拒答主观/超纲问题）：**{res['rejected_n']} 条**"
             f"（{fmt(res['rejected_pct'],1)}%）")
    g = res["guess"]
    tot = sum(g.values()) or 1
    L.append("- 同学对回答来源的猜测：")
    for k, v in g.most_common():
        L.append(f"  - {k}：{v} 条（{fmt(100*v/tot,1)}%）")

    L.append("\n## 四、按收集批次分组\n")
    L.append("| 批次 | 条数 | 同学数 | 三维度均分 |")
    L.append("|---|---|---|---|")
    for s, d in res["by_source"].items():
        L.append(f"| {s} | {d['n']} | {d['students']} | {fmt(d['mean'])} |")

    L.append("\n## 五、典型评语摘录\n")
    seen = set()
    cnt = 0
    for r in rows:
        c = str(r["comment"]).strip()
        if not c or c in seen:
            continue
        seen.add(c)
        # 挑有信息量的
        if len(c) < 8:
            continue
        L.append(f"- 「{c}」——{r['student']}（{r['major']}）")
        cnt += 1
        if cnt >= 12:
            break

    L.append("\n## 六、回填位置\n")
    L.append("1. `01-模型与算法说明文档.md` 6.6「评测结果记录表」实验三行")
    L.append("2. `07-同学盲评问卷.md` 3.1 统计表")
    L.append("3. `06-答辩PPT大纲.md` 第 9 页「评测结果」")
    L.append("4. `code/eval/eval_report.md` 新增「实验三：同学盲评」章节\n")
    return "\n".join(L)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True, help="盲评 Excel 路径")
    ap.add_argument("--out", default=None, help="输出 blind.csv 路径，默认 eval/blind.csv")
    args = ap.parse_args()

    xlsx = Path(args.xlsx)
    if not xlsx.exists():
        print("找不到文件：", xlsx, file=sys.stderr)
        sys.exit(1)

    rows = normalize(load_rows(xlsx))
    if not rows:
        print("未读到数据", file=sys.stderr)
        sys.exit(1)

    here = Path(__file__).resolve().parent
    out_csv = Path(args.out) if args.out else here / "blind.csv"
    out_csv.parent.mkdir(parents=True, exist_ok=True)

    with open(out_csv, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)

    res = stat(rows)
    md = build_md(rows, res)
    out_md = here / "盲评统计结果.md"
    out_md.write_text(md, encoding="utf-8")

    print(f"已写入 {out_csv}（{len(rows)} 条）")
    print(f"已写入 {out_md}")
    print()
    print(f"综合均分 {fmt(res['overall'])} / 5，样本 {res['n']} 条，"
          f"专业 {res['majors']} 个，同学编号 {res['students']} 个")
    for k in ["acc", "use", "trust"]:
        if k in res:
            d = res[k]
            print(f"  {d['label']}: {fmt(d['mean'])} （满分占比 {fmt(d['full_pct'],1)}%）")


if __name__ == "__main__":
    main()
